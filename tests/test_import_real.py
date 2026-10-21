r"""Optional regression against the exact user workbook; no source file bundled.

PowerShell: $env:TKD_SOURCE_XLSX='D:\TKD\<workbook>.xlsx';
python -m unittest discover -s tests -p test_import_real.py -v
"""
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import unittest

import openpyxl

from app import core, importer


SOURCE_PATH = os.environ.get("TKD_SOURCE_XLSX")


class MedalSourceInterpretationTests(unittest.TestCase):
    def test_conflicting_singleton_is_blocked(self):
        self.assertTrue(importer._medal_conflict("🥇 Nguyễn A – HCB quyền cá nhân", "HCB"))

    def test_multiple_medals_not_falsely_assigned_to_one_result(self):
        self.assertFalse(importer._medal_conflict("A HCV 🥇; B nhận 🥈", "HCB"))
        self.assertFalse(importer._medal_conflict("A HCV 🥇; B HCB 🥈", "HCV"))
        self.assertTrue(importer._multiple_medals_in_source("A HCV 🥇; B nhận 🥈"))


@unittest.skipUnless(SOURCE_PATH, "Đặt TKD_SOURCE_XLSX để kiểm tra file Excel nguồn thật.")
class RealWorkbookImportTests(unittest.TestCase):
    def test_exact_source_snapshot_pending_and_idempotent(self):
        source = Path(SOURCE_PATH)
        data = source.read_bytes()
        conn = sqlite3.connect(":memory:")
        conn.row_factory = sqlite3.Row
        self.addCleanup(conn.close)
        core.init_db(conn)
        preview = importer.preview(conn, data, source.name)
        self.assertFalse(preview["errors"])
        self.assertEqual(preview["counts"], {"athletes": 170, "tournaments": 28, "achievements": 406, "links": 448, "queue": 53})
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM athletes").fetchone()[0], 0)
        with conn:
            outcome = importer.commit_import(conn, data, source.name, "source-audit")
        self.assertTrue(outcome["imported"])
        self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM source_records").fetchone()[0], 1105)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM branches").fetchone()[0], 0)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM achievements WHERE verification_status!='pending' OR publication_allowed!=0").fetchone()[0], 0)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM verification_queue WHERE resolution_status!='unresolved'").fetchone()[0], 0)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM athletes WHERE official_code IS NOT NULL OR identity_verification_status!='pending' OR date_of_birth IS NOT NULL OR current_branch_id IS NOT NULL").fetchone()[0], 0)
        raw = json.loads(conn.execute("SELECT original_json FROM source_records WHERE sheet='03_THANH_TICH' AND excel_row=2").fetchone()[0])
        self.assertEqual(raw["raw_cells"][15], "Đã xác minh")
        self.assertTrue(raw["raw_cells"][2].startswith("=IF("))
        snapshot = json.loads(conn.execute("SELECT workbook_json FROM import_batches").fetchone()[0])
        workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
        self.addCleanup(workbook.close)
        self.assertEqual([s["name"] for s in snapshot["sheets"]], workbook.sheetnames)
        self.assertEqual(sum(len(s["rows"]) for s in snapshot["sheets"]), 5398)
        for original_sheet, saved_sheet in zip(workbook, snapshot["sheets"]):
            self.assertEqual(saved_sheet["max_row"], original_sheet.max_row)
            self.assertEqual(saved_sheet["max_column"], original_sheet.max_column)
            for number, (original, saved) in enumerate(zip(original_sheet.values, saved_sheet["rows"]), 1):
                self.assertEqual(saved["row_number"], number)
                self.assertEqual(saved["raw_cells"], [importer._raw(v) for v in original])
        with conn:
            repeated = importer.commit_import(conn, data, source.name, "source-audit")
        self.assertFalse(repeated["imported"])
        self.assertTrue(repeated["existing_file"])
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM achievements").fetchone()[0], 406)
        self.assertEqual(conn.execute("SELECT COUNT(*) FROM import_batches").fetchone()[0], 1)
        changed = importer.preview(conn, data + b"\n", "changed.xlsx")
        self.assertFalse(changed["existing_file"])
        self.assertEqual(sum(x["code"] == "LEGACY_CONFLICT" for x in changed["errors"]), 657)
        dashboard = core.Service(conn, {"id": "source-audit", "role": "admin", "can_verify": 1}).dashboard()
        self.assertEqual(dashboard["verified"], 0)
        self.assertEqual(dashboard["pending"], 406)
        self.assertEqual(dashboard["club_medals"], {"HCV": 0, "HCB": 0, "HCĐ": 0})
        self.assertEqual(dashboard["athlete_medal_credits"], 0)


if __name__ == "__main__":
    unittest.main()
