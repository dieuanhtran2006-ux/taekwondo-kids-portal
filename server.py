"""Local development server and production WSGI application factory.

Run: python server.py --import source.xlsx --port 8765
Production: waitress-serve --call server:create_application
"""
from __future__ import annotations
import argparse
import hashlib
import io
import json
import logging
import os
import sys
import secrets
import sqlite3
import threading
import time
from collections import defaultdict
from pathlib import Path
from urllib.parse import parse_qs, unquote
from wsgiref.simple_server import make_server, WSGIServer, WSGIRequestHandler
from socketserver import ThreadingMixIn
from app.core import Service, AppError, connect, init_db, now, new_id, audit, create_user, safe_user, hash_password, verify_password, permission_value, MEDALS, DISCIPLINES, PARTICIPATION_TYPES
from app.public_portal import PublicPortal

BASE=Path(__file__).resolve().parent
LOG=logging.getLogger('tkd')
MAX_BODY=15*1024*1024

class Application:
    def __init__(self,db_path=None):
        self.db_path=Path(db_path or os.environ.get('TKD_DB',BASE/'data'/'portal.sqlite3')).resolve()
        self.db_path.parent.mkdir(parents=True,exist_ok=True)
        self.storage=self.db_path.parent/'evidence'
        self.storage.mkdir(exist_ok=True)
        initial_conn=connect(self.db_path)
        try: init_db(initial_conn)
        finally: initial_conn.close()
        self.previews={}
        self.login_attempts=defaultdict(list)
        self.public_requests={}
        self.submission_requests={}
        self.lock=threading.RLock()
        self.secure=os.environ.get('TKD_SECURE_COOKIE','0')=='1'
        self.allowed_hosts=[v.strip().lower() for v in os.environ.get('TKD_ALLOWED_HOSTS','localhost,127.0.0.1,[::1]').split(',') if v.strip()]
        self.dummy_hash=hash_password(secrets.token_urlsafe(20))

    def __call__(self,environ,start_response):
        headers=[('Cache-Control','no-store'),('X-Content-Type-Options','nosniff'),('X-Frame-Options','DENY'),('Referrer-Policy','no-referrer'),('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'none'; base-uri 'self'; form-action 'self'")]
        conn=None
        try:
            host=environ.get('HTTP_HOST','').lower()
            hostname=host.rsplit(':',1)[0] if ':' in host and not host.endswith(']') else host
            if hostname not in self.allowed_hosts and host not in self.allowed_hosts: raise AppError(400,'Tên máy chủ không được phép.')
            method=environ['REQUEST_METHOD'].upper(); path=unquote(environ.get('PATH_INFO','/'))
            if method not in ('GET','POST','PATCH','DELETE'): raise AppError(405,'Phương thức không được hỗ trợ.')
            if method!='GET':
                origin=environ.get('HTTP_ORIGIN')
                expected=environ.get('wsgi.url_scheme','http')+'://'+host
                if origin and origin!=expected: raise AppError(403,'Nguồn yêu cầu không hợp lệ.')
                if environ.get('HTTP_SEC_FETCH_SITE')=='cross-site': raise AppError(403,'Yêu cầu khác nguồn bị từ chối.')
            if not path.startswith('/api/'):
                if method!='GET': raise AppError(405,'Phương thức không hợp lệ.')
                data,mime=self.static(path)
                headers += [('Content-Type',mime)]
                start_response('200 OK',headers+[('Content-Length',str(len(data)))])
                return [data]
            conn=connect(self.db_path)
            user,session=self.session(conn,environ)
            if method!='GET' and user and path not in ('/api/login','/api/student/login'):
                token=environ.get('HTTP_X_CSRF_TOKEN','')
                if not session or not secrets.compare_digest(token,session['csrf']): raise AppError(403,'Phiên thao tác không hợp lệ. Hãy tải lại trang.')
            length=int(environ.get('CONTENT_LENGTH') or 0)
            if path in ('/api/login','/api/student/login','/api/password') and length>8192:
                raise AppError(413,'Yêu cầu đăng nhập hoặc đổi mật khẩu quá dài.')
            if path=='/api/public/submissions' and length>16*1024:
                raise AppError(413,'Đề nghị vượt giới hạn 16 KB.')
            if length<0 or length>MAX_BODY: raise AppError(413,'Tệp/yêu cầu vượt giới hạn 15 MB.')
            body=environ['wsgi.input'].read(length) if length else b''
            params={k:v[-1] for k,v in parse_qs(environ.get('QUERY_STRING','')).items()}
            payload={}
            if body and 'application/json' in environ.get('CONTENT_TYPE',''):
                try: payload=json.loads(body)
                except (ValueError,UnicodeDecodeError): raise AppError(400,'JSON không hợp lệ.')
                if not isinstance(payload,dict): raise AppError(400,'Dữ liệu yêu cầu phải là object.')
            with conn:
                result,extra=self.dispatch(conn,environ,method,path,params,payload,body,user,session)
            headers+=extra
            if isinstance(result,bytes): data=result
            else:
                data=json.dumps(result,ensure_ascii=False,default=str).encode('utf-8')
                headers.append(('Content-Type','application/json; charset=utf-8'))
            start_response('200 OK',headers+[('Content-Length',str(len(data)))])
            return [data]
        except AppError as e:
            status=e.status; result={'error':e.message}
            if e.details is not None: result['details']=e.details
        except sqlite3.IntegrityError:
            status=422; result={'error':'Trùng mã/liên kết hoặc tham chiếu tới bản ghi không tồn tại.'}
        except (ValueError,TypeError,KeyError) as e:
            LOG.warning('Invalid request: %s',type(e).__name__)
            status=422; result={'error':'Dữ liệu yêu cầu không hợp lệ. Kiểm tra các trường và thử lại.'}
        except Exception:
            LOG.exception('Request failed')
            status=500; result={'error':'Không thể hoàn tất thao tác. Vui lòng thử lại; lỗi đã được ghi nhận.'}
        finally:
            if conn: conn.close()
        data=json.dumps(result,ensure_ascii=False).encode('utf-8')
        labels={400:'Bad Request',401:'Unauthorized',403:'Forbidden',404:'Not Found',405:'Method Not Allowed',409:'Conflict',413:'Payload Too Large',422:'Unprocessable Entity',429:'Too Many Requests',500:'Internal Server Error'}
        start_response(f'{status} {labels.get(status,"Error")}',headers+[('Content-Type','application/json; charset=utf-8'),('Content-Length',str(len(data)))])
        return [data]

    def static(self,path):
        if path in ('/static/student.js','/static/student.css'):
            return (BASE/'static'/path.split('/')[-1]).read_bytes(), ('text/javascript; charset=utf-8' if path.endswith('.js') else 'text/css; charset=utf-8')
        if path in ('/dang-nhap','/dang-nhap/','/dang-nhap.html'):
            return (BASE/'static'/'student.html').read_bytes(),'text/html; charset=utf-8'
        if path=='/static/contributions.js':
            return (BASE/'static'/'contributions.js').read_bytes(),'text/javascript; charset=utf-8'
        files={'/static/app.js':('app.js','text/javascript; charset=utf-8'),'/static/style.css':('style.css','text/css; charset=utf-8'),'/static/public.js':('public.js','text/javascript; charset=utf-8'),'/static/public.css':('public.css','text/css; charset=utf-8'),'/static/logo.png':('logo.png','image/png'),'/favicon.ico':('logo.png','image/png')}
        if path in files: name,mime=files[path]
        elif path in ('/', '/tra-cuu', '/tra-cuu/') or path.startswith('/ho-so/'): name,mime='public.html','text/html; charset=utf-8'
        elif path in ('/admin', '/admin/'): name,mime='index.html','text/html; charset=utf-8'
        else: raise AppError(404,'Không tìm thấy trang.')
        target=BASE/'static'/name
        if not target.is_file(): raise AppError(404,'Không tìm thấy giao diện.')
        return target.read_bytes(),mime

    def session(self,conn,environ):
        from http.cookies import SimpleCookie
        cookie=SimpleCookie()
        try: cookie.load(environ.get('HTTP_COOKIE',''))
        except Exception: return None,None
        token=cookie.get('tkd_session')
        if not token: return None,None
        session=conn.execute('SELECT * FROM sessions WHERE token_hash=? AND expires_at>?',(hashlib.sha256(token.value.encode()).hexdigest(),time.time())).fetchone()
        if not session: return None,None
        user=conn.execute('SELECT * FROM users WHERE id=? AND active=1',(session['user_id'],)).fetchone()
        return (safe_user(user),dict(session)) if user else (None,None)

    def session_response(self,conn,user,session):
        if user:
            user=dict(user)
            user['athlete_ids']=[] if user.get('must_change_password') else [r[0] for r in conn.execute('SELECT athlete_id FROM user_athletes WHERE user_id=?',(user['id'],))]
            user['branch_ids']=[] if user.get('must_change_password') else [r[0] for r in conn.execute('SELECT branch_id FROM user_branches WHERE user_id=?',(user['id'],))]
        return {'user':user,'csrf':session['csrf'] if session else '', 'configured':bool(conn.execute('SELECT 1 FROM users LIMIT 1').fetchone())}

    def start_session(self,conn,user):
        token=secrets.token_urlsafe(32); csrf=secrets.token_urlsafe(32)
        conn.execute('DELETE FROM sessions WHERE expires_at<?',(time.time(),))
        conn.execute('INSERT INTO sessions VALUES(?,?,?,?)',(hashlib.sha256(token.encode()).hexdigest(),user['id'],csrf,time.time()+8*3600))
        cookie=f'tkd_session={token}; HttpOnly; SameSite=Strict; Path=/; Max-Age=28800'+('; Secure' if self.secure else '')
        return self.session_response(conn,user,{'csrf':csrf}),[('Set-Cookie',cookie)]

    def dispatch(self,conn,env,method,path,q,p,raw,user,session):
        if path=='/api/health' and method=='GET': return {'app':'taekwondo-kids-portal','status':'ok'},[]
        if path.startswith('/api/public/'):
            if path=='/api/public/submissions' and method=='POST':
                if 'application/json' not in env.get('CONTENT_TYPE',''):
                    raise AppError(422,'Đề nghị cần được gửi dưới dạng JSON.')
                self.submission_rate_limit(env.get('REMOTE_ADDR',''))
                from app.submissions import receive
                conn.execute('BEGIN IMMEDIATE')
                return receive(conn,p),[]
            if path=='/api/public/submissions': raise AppError(405,'Đề nghị chỉ hỗ trợ gửi mới; danh sách được bảo vệ trong khu quản trị.')
            if method!='GET': raise AppError(405,'Trang công khai chỉ hỗ trợ xem dữ liệu.')
            self.public_rate_limit(env.get('REMOTE_ADDR', ''))
            conn.execute('BEGIN')
            portal=PublicPortal(conn)
            if path=='/api/public/summary': return portal.summary(),[]
            if path=='/api/public/athletes': return portal.search(q),[]
            parts=path.strip('/').split('/')
            if len(parts)==4:
                if parts[2]=='athletes': return portal.profile(parts[3]),[]
                if parts[2]=='achievements': return portal.achievement(parts[3]),[]
                if parts[2]=='evidence': return self.read_evidence(portal.evidence(parts[3]))
            raise AppError(404,'Không tìm thấy chức năng công khai.')
        if path=='/api/session' and method=='GET': return self.session_response(conn,user,session),[]
        if path=='/api/setup' and method=='POST':
            if os.environ.get('TKD_ALLOW_SETUP','1')!='1' or env.get('REMOTE_ADDR') not in ('127.0.0.1','::1'):
                raise AppError(403,'Khởi tạo chỉ được phép trên máy chủ cục bộ.')
            conn.execute('BEGIN IMMEDIATE')
            if conn.execute('SELECT 1 FROM users LIMIT 1').fetchone(): raise AppError(409,'Hệ thống đã có tài khoản quản trị.')
            new_user=create_user(conn,{**p,'role':'admin','can_verify':True,'athlete_ids':[],'branch_ids':[]})
            audit(conn,new_user['id'],'setup','users',new_user['id'],None,{'username':new_user['username'],'role':'admin'})
            return self.start_session(conn,new_user)
        if path in ('/api/login','/api/student/login') and method=='POST':
            username=str(p.get('username','')).strip()
            key=(env.get('REMOTE_ADDR',''),username.casefold())
            with self.lock:
                self.login_attempts[key]=[t for t in self.login_attempts[key] if t>time.time()-900]
                if len(self.login_attempts[key])>=8: raise AppError(429,'Quá nhiều lần đăng nhập sai. Thử lại sau 15 phút.')
            row=conn.execute('SELECT * FROM users WHERE username=? AND active=1',(username,)).fetchone()
            valid=verify_password(p.get('password',''),row['password_hash'] if row else self.dummy_hash)
            if not row or not valid or (path=='/api/student/login' and row['role']!='viewer'):
                with self.lock: self.login_attempts[key].append(time.time())
                raise AppError(401,'Tên đăng nhập hoặc mật khẩu không đúng.')
            with self.lock: self.login_attempts.pop(key,None)
            audit(conn,row['id'],'login','users',row['id'])
            return self.start_session(conn,safe_user(row))
        if not user: raise AppError(401,'Vui lòng đăng nhập.')
        if user.get('must_change_password') and path not in ('/api/password','/api/logout'):
            raise AppError(403,'Cần đổi mật khẩu lần đầu trước khi mở hồ sơ.',{'code':'PASSWORD_CHANGE_REQUIRED'})
        if path=='/api/password' and method=='POST':
            conn.execute('BEGIN IMMEDIATE')
            before=conn.execute('SELECT * FROM users WHERE id=?',(user['id'],)).fetchone()
            if not verify_password(p.get('current_password',''),before['password_hash']): raise AppError(422,'Mật khẩu hiện tại chưa đúng.')
            password=p.get('password','')
            if (before['must_change_password'] or 'confirm_password' in p) and password!=p.get('confirm_password'):
                raise AppError(422,'Mật khẩu mới và nhập lại chưa khớp.')
            if verify_password(password,before['password_hash']) or (isinstance(password,str) and password.casefold()==before['username'].casefold()):
                raise AppError(422,'Mật khẩu mới phải khác mật khẩu hiện tại và mã đăng nhập.')
            conn.execute('UPDATE users SET password_hash=?,must_change_password=0,password_changed_at=?,version=version+1 WHERE id=?',(hash_password(password),now(),user['id']))
            conn.execute('DELETE FROM sessions WHERE user_id=?',(user['id'],))
            audit(conn,user['id'],'password_changed','users',user['id'])
            refreshed=safe_user(conn.execute('SELECT * FROM users WHERE id=?',(user['id'],)).fetchone())
            return self.start_session(conn,refreshed)
        if path=='/api/logout' and method=='POST':
            conn.execute('DELETE FROM sessions WHERE token_hash=?',(session['token_hash'],))
            return {'ok':True},[('Set-Cookie','tkd_session=; HttpOnly; SameSite=Strict; Path=/; Max-Age=0'+('; Secure' if self.secure else ''))]
        svc=Service(conn,user)
        if path.startswith('/api/student/athletes/') and method=='GET':
            from app.student_accounts import student_profile
            return student_profile(conn,user,path.split('/')[-1]),[]
        if path.startswith('/api/student/achievements/') and method=='GET':
            if user['role']!='viewer': raise AppError(403,'Chức năng dành cho tài khoản võ sinh.')
            from app.student_accounts import student_result
            return student_result(svc.detail('achievements',path.split('/')[-1])),[]
        if path.startswith('/api/athletes/') and path.endswith('/account') and method=='POST':
            from app.student_accounts import issue_account
            conn.execute('BEGIN IMMEDIATE')
            return issue_account(conn,user,path.split('/')[-2],p),[]
        if path=='/api/submissions' or path.startswith('/api/submissions/'):
            from app.submissions import SubmissionReview
            review=SubmissionReview(conn,user)
            parts=path.strip('/').split('/')
            if len(parts)==2 and method=='GET': return review.list(q),[]
            if len(parts)==3:
                if method=='GET': return review.detail(parts[2]),[]
                if method=='PATCH':
                    conn.execute('BEGIN IMMEDIATE')
                    return review.update(parts[2],p),[]
            if len(parts)==4 and parts[3]=='convert' and method=='POST':
                conn.execute('BEGIN IMMEDIATE')
                return review.convert(parts[2],p),[]
            raise AppError(405,'Thao tác đề nghị không được hỗ trợ.')
        if path.startswith('/api/publication/') and method=='PATCH':
            parts=path.strip('/').split('/')
            if len(parts)==4: return svc.set_publication(parts[2],parts[3],p),[]
            raise AppError(404,'Không tìm thấy chức năng công bố.')
        if path=='/api/catalogs' and method=='GET':
            return {'disciplines':DISCIPLINES,'participation_types':PARTICIPATION_TYPES,'medals':MEDALS,'statuses':['Đang tập','Tạm nghỉ','Ngừng tập','Chuyển cơ sở','Chưa xác định'],'issues':['Thiếu người nhận','Tên viết tắt','Thiếu hoặc mâu thuẫn thông tin','Chưa rõ hình thức hoặc nội dung thi','Thành tích đồng đội chưa đủ thành viên','Khác']},[]
        if path=='/api/dashboard' and method=='GET': return svc.dashboard(q),[]
        if path=='/api/reports' and method=='GET': return svc.reports(q),[]
        if path=='/api/tournament-catalog' and method=='GET':
            svc.require_staff()
            return {'items':[dict(r) for r in conn.execute('SELECT id,tournament_name,year,competition_level,status FROM tournaments WHERE archived=0 ORDER BY year DESC,tournament_name')]},[]
        if path=='/api/integrations' and method=='GET':
            svc.require_admin()
            from app.google_adapter import integration_status
            return {'google':integration_status()},[]
        if path=='/api/audit' and method=='GET':
            svc.require_admin(); page=max(1,int(q.get('page',1))); limit=max(1,min(100,int(q.get('limit',30))))
            total=conn.execute('SELECT COUNT(*) FROM audit_logs').fetchone()[0]
            rows=[dict(r) for r in conn.execute('SELECT a.*,COALESCE(u.username,a.actor) actor FROM audit_logs a LEFT JOIN users u ON u.id=a.actor ORDER BY a.id DESC LIMIT ? OFFSET ?',(limit,(page-1)*limit))]
            return {'items':rows,'total':total,'page':page,'pages':max(1,(total+limit-1)//limit)},[]
        if path.startswith('/api/source/') and method=='GET':
            svc.require_staff(); _,_,_,entity,id=path.split('/')
            svc._authorize(svc.table(entity),id)
            if svc.table(entity)=='achievements': svc.achievement_scope(id,True)
            if svc.table(entity)=='tournaments' and not svc.admin: raise AppError(403,'Nguồn giải đầy đủ chỉ dành cho quản trị viên.')
            rows=[dict(r) for r in conn.execute('SELECT s.*,b.filename,b.file_hash FROM source_records s JOIN import_batches b ON b.id=s.batch_id WHERE s.entity_id=?',(id,))]
            return {'items':rows},[]
        if path.startswith('/api/import/'):
            svc.require_admin()
            from app.importer import preview,commit_import
            action=path.rsplit('/',1)[-1]
            if action=='preview' and method=='POST':
                if not raw or not raw.startswith(b'PK'): raise AppError(422,'Chọn tệp Excel .xlsx hợp lệ.')
                filename=Path(unquote(env.get('HTTP_X_FILENAME','import.xlsx'))).name
                result=preview(conn,raw,filename); token=secrets.token_urlsafe(32)
                with self.lock:
                    self.previews={k:v for k,v in self.previews.items() if v['expires']>time.time()}
                    if len(self.previews)>=5: raise AppError(429,'Có quá nhiều lượt xem trước đang mở. Hủy lượt cũ rồi thử lại.')
                    self.previews[token]={'bytes':raw,'filename':filename,'user':user['id'],'expires':time.time()+1800}
                return {**result,'token':token},[]
            if action in ('commit','cancel') and method=='POST':
                with self.lock: item=self.previews.get(p.get('token'))
                if not item or item['user']!=user['id'] or item['expires']<time.time(): raise AppError(409,'Bản xem trước đã hết hạn. Vui lòng chọn lại tệp.')
                if action=='cancel':
                    with self.lock: self.previews.pop(p['token'],None)
                    return {'ok':True},[]
                conn.execute('BEGIN IMMEDIATE')
                result=commit_import(conn,item['bytes'],item['filename'],user['id'])
                audit(conn,user['id'],'import','import_batches',result.get('batch_id'),None,{k:v for k,v in result.items() if k!='warnings'})
                with self.lock: self.previews.pop(p['token'],None)
                return result,[]
        if path.startswith('/api/evidence'):
            return self.evidence(conn,svc,env,method,path,q,p,raw)
        if path=='/api/export' and method=='GET':
            svc.require_admin(); output=self.export(conn)
            audit(conn,user['id'],'export','reports',None)
            return output,[('Content-Type','application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'),('Content-Disposition','attachment; filename="TKD-report.xlsx"')]
        if path=='/api/users' or path.startswith('/api/users/'):
            svc.require_admin()
            if method=='GET' and path=='/api/users':
                rows=[]
                for row in conn.execute('SELECT * FROM users ORDER BY username'):
                    u=safe_user(row); u['athlete_ids']=[r[0] for r in conn.execute('SELECT athlete_id FROM user_athletes WHERE user_id=?',(u['id'],))]; u['branch_ids']=[r[0] for r in conn.execute('SELECT branch_id FROM user_branches WHERE user_id=?',(u['id'],))]; rows.append(u)
                return {'items':rows},[]
            if method=='POST' and path=='/api/users':
                u=create_user(conn,p); audit(conn,user['id'],'create','users',u['id'],None,u); return u,[]
            if method=='PATCH' and path.startswith('/api/users/'):
                conn.execute('BEGIN IMMEDIATE')
                target=path.split('/')[-1]; old=conn.execute('SELECT * FROM users WHERE id=?',(target,)).fetchone()
                if not old: raise AppError(404,'Không tìm thấy tài khoản.')
                if old['version']!=p.get('version'): raise AppError(409,'Tài khoản đã thay đổi. Tải lại trang.')
                role=p.get('role',old['role']); active=int(bool(p.get('active',old['active'])))
                if role not in ('admin','coach','viewer'): raise AppError(422,'Vai trò không hợp lệ.')
                if old['login_athlete_id'] and (role!='viewer' or ('athlete_ids' in p and set(p['athlete_ids'])!={old['login_athlete_id']}) or p.get('branch_ids')):
                    raise AppError(422,'Tài khoản theo mã võ sinh chỉ được gắn với đúng hồ sơ đã cấp.')
                if old['role']=='admin' and (role!='admin' or not active):
                    count=conn.execute("SELECT COUNT(*) FROM users WHERE role='admin' AND active=1").fetchone()[0]
                    if count<=1: raise AppError(422,'Cần giữ ít nhất một quản trị viên đang hoạt động.')
                changed=conn.execute('UPDATE users SET role=?,can_verify=?,active=?,version=version+1 WHERE id=? AND version=?',(role,int(role=='admin' or (role=='coach' and p.get('can_verify',old['can_verify']))),active,target,old['version'])).rowcount
                if changed!=1: raise AppError(409,'Tài khoản đã thay đổi. Tải lại trang.')
                if p.get('password'):
                    if str(p['password']).casefold()==old['username'].casefold(): raise AppError(422,'Mật khẩu phải khác mã đăng nhập.')
                    conn.execute('UPDATE users SET password_hash=?,must_change_password=?,password_changed_at=NULL WHERE id=?',(hash_password(p['password']),int(role=='viewer'),target))
                elif role=='viewer' and old['role']!='viewer':
                    conn.execute('UPDATE users SET must_change_password=1 WHERE id=?',(target,))
                for key,table,column in [('athlete_ids','user_athletes','athlete_id'),('branch_ids','user_branches','branch_id')]:
                    if key in p:
                        conn.execute(f'DELETE FROM {table} WHERE user_id=?',(target,))
                        conn.executemany(f'INSERT INTO {table}(user_id,{column}) VALUES(?,?)',[(target,x) for x in set(p[key])])
                conn.execute('DELETE FROM sessions WHERE user_id=?',(target,))
                u=safe_user(conn.execute('SELECT * FROM users WHERE id=?',(target,)).fetchone()); audit(conn,user['id'],'update','users',target,safe_user(old),u); return u,[]
        parts=path.strip('/').split('/')
        if len(parts)>=2 and parts[1] in ('athletes','tournaments','achievements','branches','queue'):
            entity=parts[1]
            if len(parts)==2:
                if method=='GET': return svc.list(entity,q),[]
                if method=='POST': return svc.create(entity,p),[]
            if len(parts)==3:
                if method=='GET': return svc.detail(entity,parts[2]),[]
                if method=='PATCH': return svc.update(entity,parts[2],p),[]
                if method=='DELETE': return svc.archive(entity,parts[2],p),[]
            if len(parts)==4 and entity=='achievements' and method=='POST':
                if parts[3]=='verify': return svc.verify(parts[2],p),[]
                if parts[3]=='return': return svc.return_achievement(parts[2],p),[]
        raise AppError(404,'Không tìm thấy chức năng.')

    def evidence(self,conn,svc,env,method,path,q,p,raw):
        if path=='/api/evidence' and method=='POST':
            svc.require_staff(); achievement_id=q.get('achievement_id'); before=svc.achievement_scope(achievement_id,True)
            filename=Path(unquote(env.get('HTTP_X_FILENAME',''))).name
            mime=None
            if raw.startswith(b'%PDF-'): mime='application/pdf'
            elif raw.startswith(b'\x89PNG\r\n\x1a\n'): mime='image/png'
            elif raw.startswith(b'\xff\xd8\xff'): mime='image/jpeg'
            if not mime or not filename or len(raw)>10*1024*1024: raise AppError(422,'Minh chứng chỉ nhận PDF, PNG, JPEG tối đa 10 MB.')
            id=new_id(); storage_name=id+'.bin'; target=self.storage/storage_name
            target.write_bytes(raw)
            try:
                conn.execute('INSERT INTO evidence_files(id,achievement_id,filename,storage_name,mime_type,byte_size,sha256,uploaded_by,created_at) VALUES(?,?,?,?,?,?,?,?,?)',(id,achievement_id,filename[:200],storage_name,mime,len(raw),hashlib.sha256(raw).hexdigest(),svc.user['id'],now()))
                conn.execute("UPDATE achievements SET verification_status='pending',verified_by=NULL,verified_at=NULL,public_search_allowed=0,public_evidence_allowed=0,version=version+1 WHERE id=?",(achievement_id,))
                conn.execute("UPDATE athlete_achievements SET link_status='pending' WHERE achievement_id=?",(achievement_id,))
                audit(conn,svc.user['id'],'upload_evidence','achievements',achievement_id,before,{'evidence_id':id,'filename':filename})
            except Exception:
                target.unlink(missing_ok=True); raise
            return {'id':id,'filename':filename},[]
        id=path.split('/')[-1]
        row=conn.execute('SELECT * FROM evidence_files WHERE id=?',(id,)).fetchone()
        if not row: raise AppError(404,'Không tìm thấy minh chứng.')
        svc.achievement_scope(row['achievement_id'],method!='GET' or svc.user['role']=='coach')
        if method=='GET':
            if not svc.staff and not row['publication_allowed']: raise AppError(403,'Minh chứng chưa được phép công khai cho người xem hồ sơ.')
            return self.read_evidence(row)
        if method=='PATCH':
            svc.require_admin()
            allowed=permission_value(p.get('publication_allowed',row['publication_allowed']))
            if p.get('public_access_allowed') and not allowed:
                raise AppError(422,'Cần bật quyền xem minh chứng trước khi công bố tệp cho khách.')
            public=permission_value(p.get('public_access_allowed',row['public_access_allowed'])) if allowed else 0
            conn.execute('UPDATE evidence_files SET publication_allowed=?,public_access_allowed=? WHERE id=?',(allowed,public,id))
            audit(conn,svc.user['id'],'evidence_visibility','evidence_files',id,{'publication_allowed':row['publication_allowed'],'public_access_allowed':row['public_access_allowed']},{'publication_allowed':allowed,'public_access_allowed':public})
            return {'id':id,'filename':row['filename'],'publication_allowed':allowed,'public_access_allowed':public},[]
        raise AppError(405,'Phương thức không hợp lệ.')

    def read_evidence(self,row):
        target=(self.storage/row['storage_name']).resolve()
        if target.parent!=self.storage.resolve() or not target.is_file(): raise AppError(404,'Không tìm thấy tệp minh chứng.')
        suffix={'image/png':'.png','image/jpeg':'.jpg','application/pdf':'.pdf'}[row['mime_type']]
        return target.read_bytes(),[('Content-Type',row['mime_type']),('Content-Disposition',f'attachment; filename="evidence-{row["id"]}{suffix}"')]

    def public_rate_limit(self,remote):
        moment=time.monotonic()
        with self.lock:
            if len(self.public_requests)>2048:
                self.public_requests={key:value for key,value in self.public_requests.items() if value[0]>moment-60}
            if remote not in self.public_requests and len(self.public_requests)>=4096:
                raise AppError(429,'Tra cứu đang bận. Vui lòng thử lại sau một phút.')
            start,count=self.public_requests.get(remote,(moment,0))
            if start<=moment-60: start,count=moment,0
            if count>=60: raise AppError(429,'Bạn tra cứu quá nhanh. Vui lòng thử lại sau một phút.')
            self.public_requests[remote]=(start,count+1)

    def submission_rate_limit(self,remote):
        moment=time.monotonic()
        with self.lock:
            if len(self.submission_requests)>2048:
                self.submission_requests={key:value for key,value in self.submission_requests.items() if value[0]>moment-3600}
            if remote not in self.submission_requests and len(self.submission_requests)>=4096:
                raise AppError(429,'Tiếp nhận đang bận. Vui lòng thử lại sau.')
            start,count=self.submission_requests.get(remote,(moment,0))
            if start<=moment-3600: start,count=moment,0
            if count>=5: raise AppError(429,'Bạn đã gửi nhiều đề nghị. Vui lòng thử lại sau một giờ.')
            self.submission_requests[remote]=(start,count+1)

    def export(self,conn):
        from openpyxl import Workbook
        book=Workbook(); book.remove(book.active)
        for table,title in [('athletes','Vo sinh'),('tournaments','Giai dau'),('achievements','Thanh tich'),('athlete_achievements','Lien ket'),('verification_queue','Can xac minh')]:
            sheet=book.create_sheet(title); cursor=conn.execute(f'SELECT * FROM {table}'); names=[x[0] for x in cursor.description]; sheet.append(names)
            for row in cursor:
                sheet.append([("'"+v if isinstance(v,str) and v[:1] in ('=','+','-','@') else v) for v in row])
            sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        official=conn.execute("SELECT medal_type,COUNT(*) n FROM achievements WHERE verification_status='verified' AND archived=0 GROUP BY medal_type").fetchall()
        credits=conn.execute("SELECT ac.medal_type,COUNT(*) n FROM athlete_achievements aa JOIN achievements ac ON ac.id=aa.achievement_id WHERE ac.verification_status='verified' AND ac.archived=0 GROUP BY ac.medal_type").fetchall()
        stats=book.create_sheet('Thong ke',0); stats.append(['Loại','HCV','HCB','HCĐ']); stats.append(['Huy chương CLB']+[sum(r['n'] for r in official if r['medal_type']==m) for m in MEDALS]); stats.append(['Lượt võ sinh']+[sum(r['n'] for r in credits if r['medal_type']==m) for m in MEDALS]); stats.append(['Nháp chưa tính',conn.execute("SELECT COUNT(*) FROM achievements WHERE archived=0 AND verification_status!='verified'").fetchone()[0]])
        stream=io.BytesIO(); book.save(stream); return stream.getvalue()

def create_application():
    return Application()

class ThreadedWSGIServer(ThreadingMixIn,WSGIServer):
    daemon_threads=True

class QuietHandler(WSGIRequestHandler):
    def log_message(self,format,*args):
        # Do not log query strings, raw names, source text, passwords, cookies, or file paths.
        LOG.info('%s %s',self.command,self.path.split('?')[0])

def main():
    for stream in (sys.stdout,sys.stderr):
        if hasattr(stream,'reconfigure'): stream.reconfigure(encoding='utf-8',errors='replace')
    parser=argparse.ArgumentParser(description='Taekwondo Kids Việt Nam — localhost development server')
    parser.add_argument('--db',default=str(BASE/'data'/'portal.sqlite3'))
    parser.add_argument('--port',type=int,default=8765)
    parser.add_argument('--host',default='127.0.0.1')
    parser.add_argument('--import',dest='source')
    args=parser.parse_args()
    logging.basicConfig(level=logging.INFO,format='%(asctime)s %(levelname)s %(message)s')
    if args.host not in ('127.0.0.1','::1','localhost'):
        parser.error('Dev server chỉ bind loopback. Dùng WSGI production để triển khai qua HTTPS.')
    app=Application(args.db)
    if args.source:
        from app.importer import commit_import
        path=Path(args.source)
        conn=connect(app.db_path)
        try:
            with conn:
                result=commit_import(conn,path.read_bytes(),path.name,'local-bootstrap')
                audit(conn,'local-bootstrap','import','import_batches',result.get('batch_id'),None,result)
                LOG.info('Import: %s',json.dumps({k:v for k,v in result.items() if k not in ('warnings','errors')},ensure_ascii=True))
        finally: conn.close()
    with make_server(args.host,args.port,app,server_class=ThreadedWSGIServer,handler_class=QuietHandler) as httpd:
        print(f'Taekwondo Kids Portal: http://{args.host}:{args.port}',flush=True)
        print(f'Tạo tài khoản / đăng nhập quản trị: http://{args.host}:{args.port}/admin/',flush=True)
        try: httpd.serve_forever()
        except KeyboardInterrupt: pass

if __name__=='__main__': main()
